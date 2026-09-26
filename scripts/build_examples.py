#!/usr/bin/env python3
"""Rebuild the self-authored demonstration sources; no private data or model calls."""
import argparse
import base64
from datetime import datetime, timezone
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1] / "examples/sources"
ROWS = (("Alpha", "1.25"), ("Beta", "0.75"), ("Total", "2.00"))


def pdf():
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.colors import HexColor
    canvas = Canvas(str(ROOT / "sample.pdf"), pagesize=(612, 792), invariant=1)
    canvas.setTitle("A small document, fully traceable")
    canvas.setAuthor("Knowledge Base Kit contributors")
    canvas.setSubject("Self-authored extraction and vision fixture")

    def header(number, title):
        canvas.setFillColor(HexColor("#173348"))
        canvas.setFont("Helvetica-Bold", 23)
        canvas.drawString(48, 720, title)
        canvas.setFont("Helvetica", 10)
        canvas.setFillColor(HexColor("#526574"))
        canvas.drawString(48, 747, "KNOWLEDGE BASE KIT / SELF-AUTHORED EXAMPLE")
        canvas.drawString(48, 38, "Fictional measurements for testing only")
        canvas.drawRightString(564, 38, str(number))

    header(1, "A small document, fully traceable")
    canvas.setFillColor(HexColor("#172a38"))
    canvas.setFont("Helvetica", 12)
    canvas.drawString(48, 680, "Two samples are combined. Preserve every value and its unit.")
    canvas.drawString(48, 660, "The measurements below are invented for this example.")
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(48, 607, "1. Measurements")
    table_y = 568
    for i, row in enumerate((("Sample", "Volume (L)"), *ROWS)):
        y = table_y - i * 42
        canvas.setFillColor(HexColor("#e9f0f5" if i % 2 == 0 else "#ffffff"))
        canvas.rect(48, y - 29, 516, 42, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#172a38"))
        canvas.setFont("Helvetica-Bold" if i in (0, 3) else "Helvetica", 12)
        canvas.drawString(65, y - 12, row[0])
        canvas.drawString(335, y - 12, row[1])
    canvas.setStrokeColor(HexColor("#bacbd6"))
    for i in range(5):
        canvas.line(48, table_y + 13 - i * 42, 564, table_y + 13 - i * 42)
    for x in (48, 310, 564):
        canvas.line(x, table_y + 13, x, table_y + 13 - 168)
    canvas.setFont("Helvetica-Bold", 16)
    canvas.drawString(48, 349, "2. Check the sum")
    canvas.setFont("Helvetica", 17)
    canvas.drawString(65, 310, "V_total = V_Alpha + V_Beta = 1.25 + 0.75 = 2.00 L")
    canvas.setFont("Helvetica", 12)
    canvas.drawString(48, 262, "Keep decimal precision. A zero value must remain visible.")
    canvas.drawString(48, 241, "Control sample: 0.00 L. Do not treat zero as a missing cell.")
    canvas.showPage()
    header(2, "A reviewable workflow")
    canvas.setFont("Helvetica", 12)
    canvas.setFillColor(HexColor("#172a38"))
    canvas.drawString(48, 679, "Read each arrow and both outcomes before writing a description.")

    def box(x, y, label, color):
        canvas.setFillColor(HexColor(color))
        canvas.roundRect(x, y, 142, 64, 10, stroke=0, fill=1)
        canvas.setFillColor(HexColor("#172a38"))
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawCentredString(x + 71, y + 27, label)

    def arrow(x1, y1, x2, y2, label=""):
        canvas.setStrokeColor(HexColor("#526574"))
        canvas.setFillColor(HexColor("#526574"))
        canvas.setLineWidth(2)
        canvas.line(x1, y1, x2, y2)
        p = canvas.beginPath()
        if x1 == x2:
            p.moveTo(x2, y2); p.lineTo(x2 - 5, y2 + 9); p.lineTo(x2 + 5, y2 + 9)
        else:
            p.moveTo(x2, y2); p.lineTo(x2 - 9, y2 - 5); p.lineTo(x2 - 9, y2 + 5)
        p.close(); canvas.drawPath(p, fill=1, stroke=0)
        if label:
            canvas.setFont("Helvetica", 11)
            canvas.drawString((x1+x2)/2 + (8 if x1 == x2 else -12), (y1+y2)/2 + 13, label)

    box(48, 516, "Intake", "#e9f0f5")
    box(234, 516, "Validate", "#e9f0f5")
    box(420, 516, "Publish", "#d9f0e2")
    box(234, 368, "Revise", "#fde6cb")
    arrow(190, 548, 234, 548)
    arrow(376, 548, 420, 548, "pass")
    arrow(305, 516, 305, 432, "fail")
    canvas.setFillColor(HexColor("#172a38"))
    canvas.setFont("Helvetica", 12)
    for i, line in enumerate(("Intake leads to Validate.", "A passing check leads to Publish (green).",
                              "A failed check leads to Revise (orange).", "After revision, request a fresh check; do not imply automatic approval.")):
        canvas.drawString(48, 297 - i * 23, line)
    canvas.save()
    import pymupdf
    with pymupdf.open(ROOT / "sample.pdf") as doc:
        doc[1].get_pixmap(dpi=130, clip=pymupdf.Rect(35, 210, 578, 439)).save(ROOT / "workflow.png")


def docx():
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    document = Document()
    for section in document.sections:
        section.top_margin = section.bottom_margin = Inches(0.7)
    document.styles["Normal"].font.name = "Arial"
    document.styles["Normal"].font.size = Pt(11)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    document.add_heading("Bench notebook", 0)
    document.add_paragraph("Self-authored fixture. Paragraph before the table.")
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Sample", "Volume (L)"
    for sample, volume in ROWS:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = sample, volume
    document.add_paragraph("Paragraph after the table, before the workflow image.")
    document.add_picture(str(ROOT / "workflow.png"), width=Inches(5.8))
    document.add_paragraph("Paragraph after the image. Pass leads to Publish; fail leads to Revise.")
    document.add_heading("A useful caveat", 1)
    document.add_paragraph("These are invented measurements, not a laboratory protocol. Keep 0.00 as a value.")
    props = document.core_properties
    props.author = "Knowledge Base Kit contributors"
    props.last_modified_by = "Knowledge Base Kit contributors"
    props.title = "Bench notebook"
    props.created = props.modified = datetime(2026, 1, 1, tzinfo=timezone.utc)
    document.save(ROOT / "sample.docx")


def text_sources():
    (ROOT / "quickstart.md").write_text("""# A tiny knowledge library

This is an original demonstration document for Knowledge Base Kit.

## Purpose

Keep a full, source-backed document and a shorter reading guide in a local folder.
Read the catalog first, then the summary, then the full document when needed.

## Workflow

1. Put a source in the inbox.
2. Extract and review it in a draft directory.
3. Check structure, links, math and source coverage.
4. File the reviewed document in an appropriate category.

## Example values

| Sample | Volume (L) |
| --- | --- |
| Alpha | 1.25 |
| Beta | 0.75 |
| Total | 2.00 |

The sum is $1.25 + 0.75 = 2.00$. A control value of $0.00$ is not missing data.

## Caveat

The numbers are invented for testing. A structural validator cannot prove factual
accuracy; always compare the generated document with its source.
""", encoding="utf-8")
    image = base64.b64encode((ROOT / "workflow.png").read_bytes()).decode()
    (ROOT / "sample.html").write_text(f'''<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Static workflow notes</title></head>
<body><h1>Static workflow notes</h1><p>Before the table.</p>
<table><tr><th>Sample</th><th>Volume (L)</th></tr><tr><td>Alpha</td><td>1.25</td></tr>
<tr><td>Beta</td><td>0.75</td></tr><tr><td>Control</td><td>0.00</td></tr></table>
<p>After the table, before the inline image.</p>
<img src="data:image/png;base64,{image}" alt="Intake, Validate, pass to Publish, fail to Revise">
<p>After the image. <a href="#caveat">Read the caveat.</a></p>
<h2 id="caveat">Caveat</h2><p>All measurements are invented.</p></body></html>
''', encoding="utf-8")
    notebook = {"nbformat": 4, "nbformat_minor": 5,
                "metadata": {"kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"}},
                "cells": [
                    {"cell_type": "markdown", "id": "intro", "metadata": {}, "source": ["# Saved output example\n", "This notebook must be read without executing cells."]},
                    {"cell_type": "code", "id": "sum", "metadata": {}, "execution_count": None,
                     "source": ["print(1.25 + 0.75)"],
                     "outputs": [{"output_type": "stream", "name": "stdout", "text": ["2.0\n"]}]},
                    {"cell_type": "code", "id": "figure", "metadata": {}, "execution_count": None,
                     "source": ["# A saved diagram output, intentionally not executable."],
                     "outputs": [{"output_type": "display_data", "metadata": {},
                                  "data": {"image/png": image, "text/plain": "Workflow diagram"}}]},
                ]}
    (ROOT / "sample.ipynb").write_text(json.dumps(notebook, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--only", choices=("pdf", "docx", "text", "all"), default="all")
    args = parser.parse_args()
    ROOT.mkdir(parents=True, exist_ok=True)
    for name, function in (("pdf", pdf), ("docx", docx), ("text", text_sources)):
        if args.only in (name, "all"):
            function()
            print(f"Built {name} examples")
